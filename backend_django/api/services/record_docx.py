"""The AI Use Record as an accessible Word document.

Headings use Word's Heading styles (so they are tagged headings and show in
the navigation pane), the language is en-US, the title is in the document
properties, and the logo has alt text. Same blocks and order as the PDF.
"""
import io
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from api.services.record_document import HEADING, ITEM, LOGO_ALT_TEXT, SUBHEADING, TITLE, Block

LOGO_PATH = Path(__file__).resolve().parent.parent / 'assets' / 'raise-mark.png'
LANGUAGE = 'en-US'
HEADING_LEVEL = {TITLE: 1, HEADING: 2, SUBHEADING: 3}


def _set_language(document: Document) -> None:
    """Mark the default text language as en-US for spell check and screen readers."""
    rpr_default = document.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault'))
    rpr = rpr_default.find(qn('w:rPr'))
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), LANGUAGE)
    document.core_properties.language = LANGUAGE


def _add_logo(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(str(LOGO_PATH), width=Inches(0.45))
    doc_pr = run._r.find('.//' + qn('wp:docPr'))
    doc_pr.set('descr', LOGO_ALT_TEXT)
    doc_pr.set('title', 'RAISE logo')
    name = paragraph.add_run('  RAISE · Responsible AI Standards & Ethics')
    name.bold = True
    name.font.color.rgb = RGBColor(0x00, 0x67, 0x47)


def record_docx(blocks: list[Block], title: str) -> bytes:
    document = Document()
    document.core_properties.title = title
    document.core_properties.subject = 'AI Use Record exported from RAISE, a research prototype'
    document.core_properties.author = 'RAISE'
    _set_language(document)
    document.styles['Normal'].font.size = Pt(10.5)
    _add_logo(document)

    for block in blocks:
        if block.kind in HEADING_LEVEL:
            document.add_heading(block.text, level=HEADING_LEVEL[block.kind])
        elif block.kind == ITEM:
            document.add_paragraph().add_run(block.text).bold = True
            for line in block.lines:
                document.add_paragraph(line, style='List Bullet')
        else:
            document.add_paragraph(block.text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
