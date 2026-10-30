"""Word (.docx) version of a course AI statement, with real heading styles.

The saved text is plain: its first line is the title (Heading 1), and lines
that match a section title (Heading 2) start sections. Everything else is a
normal paragraph. The student version follows on its own page, titled with
Heading 1, so screen readers and Word's navigation pane see the structure.
"""
import io

from docx import Document
from docx.enum.text import WD_BREAK

from api.services.course_statement import SECTION_HEADINGS, STUDENT_VERSION_TITLE

DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def _add_text(document: Document, text: str, title_level: int) -> None:
    lines = [line.strip() for line in text.splitlines()]
    first = True
    for line in lines:
        if not line:
            continue
        if first:
            document.add_heading(line, level=title_level)
            first = False
        elif line in SECTION_HEADINGS:
            document.add_heading(line, level=2)
        else:
            document.add_paragraph(line)


def statement_docx(statement_text: str, student_text: str, course_name: str) -> bytes:
    document = Document()
    document.core_properties.title = f'AI in {course_name}'
    _add_text(document, statement_text, title_level=1)
    if student_text.strip():
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        body = student_text.strip()
        if not body.startswith(STUDENT_VERSION_TITLE):
            body = f'{STUDENT_VERSION_TITLE}\n{body}'
        _add_text(document, body, title_level=1)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
