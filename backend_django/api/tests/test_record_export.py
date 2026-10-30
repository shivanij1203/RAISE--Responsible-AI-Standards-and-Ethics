"""The accessible AI Use Record: real text in reading order, metadata, and Word headings."""
import io
import re

from django.contrib.auth.models import User
from django.test import Client, TestCase
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from api.models import UserProfile
from api.services.record_document import LOGO_ALT_TEXT, RECORD_DISCLAIMER

DISCLOSURE = 'DISCLOSURE OF TECHNOLOGY USE\n\nActivity: Essay feedback\nTechnology tools used: ChatGPT.'


def _flat(text: str) -> str:
    return re.sub(r'\s+', ' ', text).strip()


class RecordExportTest(TestCase):
    def setUp(self) -> None:
        self.user = User.objects.create_user(username='r@usf.edu', email='r@usf.edu', password='pw12345678')
        UserProfile.objects.create(user=self.user, role='faculty')
        self.client = Client()
        self.client.login(username='r@usf.edu', password='pw12345678')
        project = self.client.post('/api/projects', data={
            'name': 'Essay feedback', 'ai_use_case': 'grading',
            'description': 'Using ChatGPT to give first-pass feedback on essays.',
            'risk_context': {'involves_student_data': True},
        }, content_type='application/json').json()
        self.project_id = project['id']
        self.client.post(f'/api/projects/{self.project_id}/decisions', data={
            'checkpoint': 'ferpa_compliance', 'description': 'Student names are removed before any AI use.',
        }, content_type='application/json')
        # What the screen shows, in the order it shows it.
        self.screen = self.client.get(f'/api/projects/{self.project_id}').json()

    def _export(self, fmt: str, **extra) -> bytes:
        response = self.client.post(f'/api/projects/{self.project_id}/record', data={
            'format': fmt, 'scope': 'full', 'disclosure': DISCLOSURE, 'time_zone': 'America/New_York', **extra,
        }, content_type='application/json')
        self.assertEqual(response.status_code, 200, response.content[:200])
        return response.content

    def _screen_order(self) -> list[str]:
        """Checklist labels grouped by category in first-appearance order, as the dashboard lists them."""
        checkpoints = self.screen['checkpoints']
        categories = list(dict.fromkeys(c['category'] for c in checkpoints))
        return [c['label'] for cat in categories for c in checkpoints if c['category'] == cat]

    def _assert_in_order(self, text: str, expected: list[str]) -> None:
        position = 0
        for piece in expected:
            found = text.find(piece, position)
            self.assertNotEqual(found, -1, f'Missing or out of order: {piece!r}')
            position = found + len(piece)

    def test_pdf_text_is_real_and_matches_the_screen_in_order(self) -> None:
        reader = PdfReader(io.BytesIO(self._export('pdf')))
        text = _flat(' '.join(page.extract_text() for page in reader.pages))

        self._assert_in_order(text, [
            'AI Use Record', 'Essay feedback', _flat(RECORD_DISCLAIMER),
            '1. About this AI use', 'Kind of use: Grading & Assessment', 'Decisions recorded: 1',
            '2. Checklist',
            *self._screen_order(),
            '3. Decisions and sources', 'Decisions', 'Student names are removed before any AI use.',
            'Sources', 'FERPA', 'https://studentprivacy.ed.gov/ferpa',
            '4. Disclosure statement', 'DISCLOSURE OF TECHNOLOGY USE', 'Technology tools used: ChatGPT.',
        ])
        self.assertIn('Status: Addressed on', text)
        self.assertIn('Status: Open', text)

    def test_pdf_has_title_language_and_heading_bookmarks(self) -> None:
        reader = PdfReader(io.BytesIO(self._export('pdf')))

        self.assertEqual(reader.metadata.title, 'AI Use Record: Essay feedback')
        self.assertEqual(reader.trailer['/Root']['/Lang'], 'en-US')
        self.assertTrue(reader.trailer['/Root']['/ViewerPreferences']['/DisplayDocTitle'])
        outline = [item.title for item in reader.outline if not isinstance(item, list)]
        self.assertEqual(outline[0], 'AI Use Record')

    def test_pdf_has_no_unlabeled_images(self) -> None:
        reader = PdfReader(io.BytesIO(self._export('pdf')))

        self.assertEqual(sum(len(page.images) for page in reader.pages), 0)

    def test_word_version_has_real_headings_language_and_logo_alt_text(self) -> None:
        document = Document(io.BytesIO(self._export('docx')))
        headings = [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith('Heading')]

        self.assertEqual(headings[0], ('Heading 1', 'AI Use Record'))
        self.assertIn(('Heading 2', '1. About this AI use'), headings)
        self.assertIn(('Heading 2', '2. Checklist'), headings)
        self.assertIn(('Heading 2', '3. Decisions and sources'), headings)
        self.assertIn(('Heading 2', '4. Disclosure statement'), headings)
        self.assertTrue(any(style == 'Heading 3' for style, _ in headings))
        self.assertEqual(document.core_properties.title, 'AI Use Record: Essay feedback')
        self.assertEqual(document.core_properties.language, 'en-US')
        lang = document.styles.element.find(qn('w:docDefaults')).find('.//' + qn('w:lang'))
        self.assertEqual(lang.get(qn('w:val')), 'en-US')
        doc_pr = document.element.body.find('.//' + qn('wp:docPr'))
        self.assertEqual(doc_pr.get('descr'), LOGO_ALT_TEXT)

    def test_word_version_keeps_the_same_order(self) -> None:
        document = Document(io.BytesIO(self._export('docx')))
        text = _flat(' '.join(p.text for p in document.paragraphs))

        self._assert_in_order(text, ['AI Use Record', '1. About this AI use', '2. Checklist',
                                     *self._screen_order(), '3. Decisions and sources', '4. Disclosure statement'])

    def test_my_items_scope_lists_only_my_items(self) -> None:
        text = _flat(' '.join(p.extract_text() for p in PdfReader(io.BytesIO(self._export('pdf', scope='mine'))).pages))

        self.assertIn('Checklist items assigned to Faculty / PI', text)
        for checkpoint in self.screen['checkpoints']:
            if checkpoint['assignedTo'] != 'pi':
                self.assertNotIn(checkpoint['label'], text)

    def test_bad_format_is_rejected(self) -> None:
        response = self.client.post(f'/api/projects/{self.project_id}/record', data={'format': 'exe'},
                                    content_type='application/json')

        self.assertEqual(response.status_code, 400)

    def test_another_user_cannot_export(self) -> None:
        other = User.objects.create_user(username='o@usf.edu', email='o@usf.edu', password='pw12345678')
        UserProfile.objects.create(user=other, role='faculty')
        client = Client()
        client.login(username='o@usf.edu', password='pw12345678')

        response = client.post(f'/api/projects/{self.project_id}/record', data={'format': 'pdf'},
                               content_type='application/json')

        self.assertEqual(response.status_code, 404)
