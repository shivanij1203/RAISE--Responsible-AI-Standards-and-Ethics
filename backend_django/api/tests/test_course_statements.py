"""Course AI statements: the text, the BOG note, the API, the Word file, and demo purging."""
import io
from datetime import timedelta
from unittest import mock

from django.contrib.auth.models import User
from django.test import Client, TestCase
from django.utils import timezone
from docx import Document

from api.models import AITool, Checkpoint, CourseStatement, Project, UsageEvent, UserProfile
from api.rules.usf import SYLLABUS_SAMPLES
from api.services import course_statement as cs
from api.services.course_statement import StatementChoices, bog_note, build_statement, build_student_version
from api.services.demo_accounts import create_demo_user, purge_expired_demo_users


def _user(email: str) -> User:
    user = User.objects.create_user(username=email, email=email, password='pw12345678')
    UserProfile.objects.create(user=user, role='faculty')
    return user


def _project(user: User, description: str, use_case: str, tools=(), done=()) -> Project:
    project = Project.objects.create(user=user, name=description[:40], description=description,
                                     ai_use_case=use_case, category='teaching')
    for name in tools:
        project.ai_tools.add(AITool.objects.get_or_create(name=name, defaults={'vendor': 'x'})[0])
    for checkpoint_id in done:
        Checkpoint.objects.create(project=project, checkpoint_id=checkpoint_id, label=checkpoint_id,
                                  category='x', assigned_to='pi', completed=True)
    return project


class StatementTextTest(TestCase):
    def setUp(self) -> None:
        self.user = _user('p@usf.edu')
        self.grading = _project(self.user, "I'm using ChatGPT to give first-pass feedback on student essays.",
                                'grading', tools=['ChatGPT'], done=['grading_validation', 'data_deidentified'])
        self.teaching = _project(self.user, 'Lecture slides for week 3', 'teaching', tools=['Claude'])

    def _choices(self, **extra) -> StatementChoices:
        return StatementChoices(course_name='PSY 3213', term='Spring 2027',
                                projects=[self.grading, self.teaching], **extra)

    def test_statement_includes_every_selected_ai_use(self) -> None:
        text = build_statement(self._choices())

        self.assertIn('I use ChatGPT to give first-pass feedback on student essays.', text)
        self.assertIn('I use Claude to help prepare course materials.', text)

    def test_commitments_come_only_from_completed_items(self) -> None:
        text = build_statement(self._choices())

        self.assertIn('I check AI-suggested grades against my own judgment.', text)
        self.assertIn('I do not enter student names or IDs into AI tools.', text)
        self.assertNotIn('fairness', text)

    def test_provost_sample_is_used_word_for_word_with_its_source(self) -> None:
        text = build_statement(self._choices(policy='restricted'))
        sample = SYLLABUS_SAMPLES['restricted']

        for paragraph in sample.paragraphs:
            self.assertIn(paragraph, text)
        self.assertIn(sample.source.url, text)

    def test_own_policy_lists_allowed_and_not_allowed_uses(self) -> None:
        text = build_statement(self._choices(policy='custom', allowed=['brainstorming', 'grammar'],
                                             not_allowed=['drafting'], cite_line='Add a note under your name.'))

        self.assertIn('You may use AI for brainstorming and exploring ideas and grammar and spelling help.', text)
        self.assertIn('You may not use AI for drafting text you submit.', text)
        self.assertIn('How to cite or disclose AI use: Add a note under your name.', text)

    def test_student_version_is_short_and_names_each_task(self) -> None:
        text = build_student_version(self._choices(policy='restricted'))

        self.assertTrue(text.startswith('How AI is used in this course'))
        self.assertIn('give first-pass feedback on student essays', text)
        self.assertIn('help prepare course materials', text)
        self.assertIn('You may not use AI tools for any course work', text)

    def test_later_recorded_sentences_stay_in_the_statement_but_not_the_student_version(self) -> None:
        self.grading.description = ("I'm using ChatGPT to give first-pass feedback on student essays. "
                                    'Student names are removed first.')
        self.grading.save()

        self.assertIn('I use ChatGPT to give first-pass feedback on student essays. Student names are removed '
                      'first.', build_statement(self._choices()))
        self.assertNotIn('Student names are removed first', build_student_version(self._choices()))

    def test_no_use_selected_leaves_a_placeholder(self) -> None:
        text = build_statement(StatementChoices(course_name='X', term='Y'))

        self.assertIn(cs.NO_FACULTY_USE, text)

    def test_statement_never_claims_to_meet_a_requirement(self) -> None:
        for policy in ('acceptable', 'restricted', 'custom'):
            text = (build_statement(self._choices(policy=policy))
                    + build_student_version(self._choices(policy=policy))).lower()
            for phrase in ('complies', 'compliant', 'satisfies', 'meets the requirement', 'legally required'):
                self.assertNotIn(phrase, text)


class BogNoteTest(TestCase):
    def test_note_reads_pending_while_the_rule_is_proposed(self) -> None:
        with mock.patch.object(cs, 'BOG_1001_STATUS', 'proposed'):
            note = bog_note()

        self.assertEqual(note['status'], 'pending')
        self.assertIn('has proposed requiring AI disclosure in syllabi. Status: pending.', note['text'])

    def test_note_reads_pending_for_a_pending_constant(self) -> None:
        with mock.patch.object(cs, 'BOG_1001_STATUS', 'pending'):
            self.assertEqual(bog_note()['status'], 'pending')

    def test_note_reads_adopted_only_when_the_constant_says_adopted(self) -> None:
        with mock.patch.object(cs, 'BOG_1001_STATUS', 'adopted'):
            note = bog_note()

        self.assertEqual(note['status'], 'adopted')
        self.assertIn('Status: adopted', note['text'])
        self.assertTrue(note['sources'][0]['url'].startswith('https://www.flbog.edu/'))


class CourseStatementApiTest(TestCase):
    def setUp(self) -> None:
        self.user = _user('a@usf.edu')
        self.project = _project(self.user, "I'm using ChatGPT to give first-pass feedback on student essays.",
                                'grading', tools=['ChatGPT'])
        self.client = Client()
        self.client.login(username='a@usf.edu', password='pw12345678')

    def _payload(self, **extra) -> dict:
        return {'course_name': 'PSY 3213', 'term': 'Spring 2027', 'project_ids': [self.project.id],
                'policy': 'acceptable', 'allowed': ['brainstorming'], 'not_allowed': [], **extra}

    def _create(self, **extra) -> dict:
        response = self.client.post('/api/course-statements', data=self._payload(**extra),
                                    content_type='application/json')
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()

    def test_options_suggest_teaching_and_grading_uses(self) -> None:
        data = self.client.get('/api/course-statements/options').json()

        self.assertEqual(data['aiUses'][0]['id'], self.project.id)
        self.assertTrue(data['aiUses'][0]['suggested'])
        self.assertEqual({p['key'] for p in data['policies']}, {'acceptable', 'restricted'})
        self.assertEqual(data['bogNote']['status'], 'pending')

    def test_draft_saves_nothing(self) -> None:
        response = self.client.post('/api/course-statements/draft', data=self._payload(),
                                    content_type='application/json')

        self.assertEqual(response.status_code, 200)
        self.assertIn('I use ChatGPT', response.json()['statementText'])
        self.assertEqual(CourseStatement.objects.count(), 0)

    def test_create_saves_and_counts_the_event(self) -> None:
        saved = self._create()

        self.assertEqual(saved['projectIds'], [self.project.id])
        self.assertIn('Acceptable Use of Generative AI Tools', saved['statementText'])
        self.assertEqual(list(UsageEvent.objects.values_list('event_type', flat=True)),
                         [UsageEvent.COURSE_STATEMENT_CREATED])

    def test_edits_are_saved_and_reopened(self) -> None:
        saved = self._create()
        self.client.put(f"/api/course-statements/{saved['id']}", data={'statement_text': 'Edited text'},
                        content_type='application/json')

        reopened = self.client.get(f"/api/course-statements/{saved['id']}").json()

        self.assertEqual(reopened['statementText'], 'Edited text')
        self.assertEqual(reopened['courseName'], 'PSY 3213')

    def test_another_users_ai_use_cannot_be_added(self) -> None:
        other = _project(_user('b@usf.edu'), 'Secret', 'grading')
        response = self.client.post('/api/course-statements', data=self._payload(project_ids=[other.id]),
                                    content_type='application/json')

        self.assertEqual(response.status_code, 400)

    def test_another_user_cannot_open_a_statement(self) -> None:
        saved = self._create()
        _user('c@usf.edu')
        other = Client()
        other.login(username='c@usf.edu', password='pw12345678')

        self.assertEqual(other.get(f"/api/course-statements/{saved['id']}").status_code, 404)

    def test_a_use_cannot_be_both_allowed_and_not_allowed(self) -> None:
        response = self.client.post('/api/course-statements',
                                    data=self._payload(allowed=['coding'], not_allowed=['coding']),
                                    content_type='application/json')

        self.assertEqual(response.status_code, 400)

    def test_course_name_is_required(self) -> None:
        response = self.client.post('/api/course-statements', data=self._payload(course_name=' '),
                                    content_type='application/json')

        self.assertEqual(response.status_code, 400)

    def test_word_download_uses_real_headings(self) -> None:
        saved = self._create()

        response = self.client.get(f"/api/course-statements/{saved['id']}/docx")
        document = Document(io.BytesIO(response.content))
        headings = [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith('Heading')]

        self.assertEqual(response.status_code, 200)
        self.assertIn('PSY 3213 AI statement.docx', response['Content-Disposition'])
        self.assertEqual(headings, [
            ('Heading 1', 'AI in PSY 3213 (Spring 2027)'),
            ('Heading 2', 'How I use AI in this course'),
            ('Heading 2', 'How you may use AI in this course'),
            ('Heading 1', 'How AI is used in this course'),
        ])


class DemoPurgeTest(TestCase):
    def test_demo_statements_are_purged_with_the_session(self) -> None:
        demo = create_demo_user()
        statement = CourseStatement.objects.create(user=demo, course_name='Demo 101', term='Fall')
        statement.projects.set(demo.projects.all()[:1])

        purge_expired_demo_users(now=timezone.now() + timedelta(hours=25))

        self.assertFalse(CourseStatement.objects.filter(id=statement.id).exists())
        self.assertFalse(User.objects.filter(id=demo.id).exists())

    def test_demo_users_can_create_statements(self) -> None:
        client = Client()
        token = client.post('/api/auth/demo').json()['token']
        response = client.post('/api/course-statements',
                               data={'course_name': 'Demo 101', 'term': 'Fall', 'project_ids': [],
                                     'policy': 'custom'},
                               content_type='application/json', HTTP_AUTHORIZATION=f'Token {token}')

        self.assertEqual(response.status_code, 201)
